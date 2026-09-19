import io

import openpyxl
import pandas as pd
import pdfplumber
import pytest

from loan_processing.document_simulation.bank_statement_xlsx import render_bank_statement
from loan_processing.document_simulation.salary_slip_pdf import render_salary_slip


@pytest.fixture
def applicant() -> pd.Series:
    return pd.Series(
        {
            "Applicant_ID": "APP00001",
            "Full_Name": "Jane Doe",
            "Employer_Name": "Acme Corp",
            "Age": 34,
            "Employment_Type": "Salaried",
            "Monthly_Net_Income": 75000.0,
            "Gross_Income": 96153.85,
            "Total_Deductions": 21153.85,
            "Total_Existing_EMIs": 12000.0,
            "CIBIL_Score": 720,
            "Requested_Loan_Amount": 500000.0,
            "Requested_Tenure_Months": 36,
            "Average_Monthly_Bank_Balance": 150000.0,
            "Number_of_Bounced_Transactions_Last_6M": 0,
            "Cash_Flow_Month_1": 145000.0,
            "Cash_Flow_Month_2": 147000.0,
            "Cash_Flow_Month_3": 149000.0,
            "Cash_Flow_Month_4": 151000.0,
            "Cash_Flow_Month_5": 153000.0,
            "Cash_Flow_Month_6": 155000.0,
        }
    )


def test_salary_slip_contains_key_applicant_fields(applicant):
    pdf_bytes = render_salary_slip(applicant)

    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)

    assert applicant["Applicant_ID"] in text
    assert str(int(applicant["Age"])) in text
    assert applicant["Employment_Type"] in text
    assert applicant["Employer_Name"] in text
    assert "Tax Deductions" in text
    assert f"{applicant['Monthly_Net_Income']:,.2f}" in text
    assert f"{applicant['Gross_Income']:,.2f}" in text
    assert f"{applicant['Total_Deductions']:,.2f}" in text


def test_bank_statement_has_six_months_and_emi_debits(applicant):
    xlsx_bytes = render_bank_statement(applicant, seed=1)

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb["Statement"]

    assert ws["B5"].value == applicant["Applicant_ID"]

    rows = list(ws.iter_rows(min_row=19, values_only=True))
    emi_rows = [r for r in rows if r[1] == "ACH DEBIT - EMI LOAN REPAYMENT"]
    salary_rows = [r for r in rows if "SALARY" in (r[1] or "")]

    assert len(emi_rows) == 6
    assert len(salary_rows) == 6
    assert all(r[2] == pytest.approx(applicant["Total_Existing_EMIs"]) for r in emi_rows)
    assert all(r[3] == pytest.approx(applicant["Monthly_Net_Income"]) for r in salary_rows)


def test_bank_statement_realized_average_is_close_to_target(applicant):
    xlsx_bytes = render_bank_statement(applicant, seed=2)

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb["Statement"]
    realized_average = ws["B8"].value

    target = applicant["Average_Monthly_Bank_Balance"]
    assert abs(realized_average - target) / target < 0.02


def test_bank_statement_has_monthly_cash_flow_trend_section(applicant):
    xlsx_bytes = render_bank_statement(applicant, seed=2)

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb["Statement"]

    assert ws["A10"].value == "Monthly Cash Flow Trend"
    for i in range(1, 7):
        assert ws[f"A{10 + i}"].value == f"Month {i} Avg Balance:"
        value = ws[f"B{10 + i}"].value
        assert isinstance(value, (int, float))
        target = applicant[f"Cash_Flow_Month_{i}"]
        assert abs(value - target) / target < 0.03


def test_bank_statement_debit_and_credit_amounts_are_never_negative(applicant):
    xlsx_bytes = render_bank_statement(applicant, seed=3)

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb["Statement"]
    rows = list(ws.iter_rows(min_row=19, values_only=True))

    for _, _, debit, credit, _ in rows:
        assert debit is None or debit >= 0
        assert credit is None or credit >= 0


def test_bank_statement_renders_bounce_charge_rows_matching_applicant_count(applicant):
    bounced_applicant = applicant.copy()
    bounced_applicant["Number_of_Bounced_Transactions_Last_6M"] = 4
    xlsx_bytes = render_bank_statement(bounced_applicant, seed=4)

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb["Statement"]
    rows = list(ws.iter_rows(min_row=19, values_only=True))

    bounce_rows = [r for r in rows if r[1] == "ACH DEBIT - CHEQUE BOUNCE CHARGE"]
    assert len(bounce_rows) == 4
    assert all(r[2] > 0 for r in bounce_rows)  # a real debit amount, not a zero-value placeholder


def test_bank_statement_zero_bounces_renders_no_bounce_rows(applicant):
    xlsx_bytes = render_bank_statement(applicant, seed=5)  # applicant fixture already has 0 bounces

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb["Statement"]
    rows = list(ws.iter_rows(min_row=19, values_only=True))

    assert not [r for r in rows if r[1] == "ACH DEBIT - CHEQUE BOUNCE CHARGE"]


def test_bank_statement_realized_average_still_close_to_target_with_bounces(applicant):
    """Bounce-charge rows are extra debits within the calibration loop, same
    as the existing discretionary debits - the per-month day-weighted
    average must still land on target despite them."""
    bounced_applicant = applicant.copy()
    bounced_applicant["Number_of_Bounced_Transactions_Last_6M"] = 5
    xlsx_bytes = render_bank_statement(bounced_applicant, seed=6)

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb["Statement"]
    realized_average = ws["B8"].value

    target = bounced_applicant["Average_Monthly_Bank_Balance"]
    assert abs(realized_average - target) / target < 0.02
