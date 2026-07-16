import openpyxl

from loan_processing.models import (
    Decision,
    DecisionOutcome,
    ExtractedFields,
    DocumentType,
    LoanApplicationResult,
    ScoreResult,
    ValidationIssue,
    ValidationResult,
)
from loan_processing.output.excel_writer import write_excel


def test_write_excel_produces_expected_sheets():
    result = LoanApplicationResult(
        application_id="test-app-1",
        extracted=[
            ExtractedFields(
                doc_type=DocumentType.APPLICATION_FORM,
                source_filename="application.txt",
                fields={"applicant_name": "Jane Doe"},
            )
        ],
        validation=ValidationResult(
            is_valid=False,
            issues=[ValidationIssue(field="requested_amount", message="Missing field", severity="error")],
        ),
        score=ScoreResult(score=640, risk_tier="medium", factors={"debt_to_income_ratio": 0.3}),
        decision=Decision(outcome=DecisionOutcome.MANUAL_REVIEW, reasons=["Missing required field"]),
    )

    buffer = write_excel(result)
    workbook = openpyxl.load_workbook(buffer)

    assert workbook.sheetnames == ["Summary", "Extracted Data", "Validation Issues", "Score Factors"]
    summary_sheet = workbook["Summary"]
    assert summary_sheet.cell(row=2, column=1).value == "test-app-1"
